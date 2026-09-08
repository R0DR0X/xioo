from inventario_utils import find_sheet


def test_exact_name_is_found():
    assert find_sheet(['RESUMEN', 'TD', 'Stock Libre'], 'Stock Libre') == 'Stock Libre'


def test_uppercase_name_is_found():
    assert find_sheet(['RESUMEN', 'TD', 'STOCK LIBRE'], 'Stock Libre') == 'STOCK LIBRE'


def test_surrounding_spaces_are_ignored():
    assert find_sheet([' stock libre '], 'Stock Libre') == ' stock libre '


def test_missing_sheet_returns_none():
    assert find_sheet(['RESUMEN', 'TD'], 'Stock Libre') is None


def test_parse_resumen_inventario_extracts_planta_and_china():
    import openpyxl
    from inventario_utils import parse_resumen_inventario

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RESUMEN AL 7-9-2026"

    # Section 1: Planta
    ws.cell(5, 6, "RENTABILIDAD DEL INVENTARIO")
    ws.cell(6, 2, "Código")
    ws.cell(6, 3, "Producto")
    ws.cell(6, 4, "TM en \nStock Libre")
    ws.cell(7, 3, "POTA FILETE")
    ws.cell(7, 4, 100.5)
    ws.cell(8, 3, "POTA NUCAS")
    ws.cell(8, 4, 200.0)
    ws.cell(9, 3, "TOTAL")
    ws.cell(9, 4, 300.5)

    # Section 2: China
    ws.cell(11, 3, "TM en Almacén de Yantai Jiahong")
    ws.cell(12, 3, "Producto")
    ws.cell(12, 4, "TM en \nStock Libre")
    ws.cell(13, 3, "POTA FILETE")
    ws.cell(13, 4, 50.0)
    ws.cell(14, 3, "TOTAL")
    ws.cell(14, 4, 50.0)

    # Section 3: Total
    ws.cell(16, 3, "INVENTARIO TOTAL VALORIZADO")
    ws.cell(17, 3, "Producto")
    ws.cell(17, 4, "TM en \nStock Libre")
    ws.cell(18, 3, "POTA FILETE")
    ws.cell(18, 4, 150.5)
    ws.cell(19, 3, "POTA NUCAS")
    ws.cell(19, 4, 200.0)
    ws.cell(20, 3, "TOTAL")
    ws.cell(20, 4, 350.5)

    res = parse_resumen_inventario(wb)
    assert res is not None
    assert res['tot_planta'] == 300.5
    assert res['tot_china'] == 50.0
    assert res['tot_total'] == 350.5
    assert len(res['planta']) == 2
    assert len(res['china']) == 1
    assert len(res['total']) == 2


def test_parse_resumen_inventario_returns_none_without_sheet():
    import openpyxl
    from inventario_utils import parse_resumen_inventario

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    assert parse_resumen_inventario(wb) is None

